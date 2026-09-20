"""Build a small human audit workbook from aspect silver disagreements."""

from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from src.common.csv_safety import escape_dataframe_for_spreadsheet
from openpyxl.worksheet.datavalidation import DataValidation

from src.ingestion.dataset_manifest import sha256_file


ASPECT_NAMES_RU = {
    "scent": "запах",
    "softness": "мягкость",
    "moisture_and_hydration": "увлажнение",
    "detangling": "распутывание",
    "product_texture_consistency": "текстура и консистенция",
    "packaging": "упаковка",
    "hair_type_compatibility": "совместимость с типом волос",
    "hair_weight": "утяжеление волос",
    "shine_appearance": "блеск и внешний вид",
    "frizz_control": "контроль пушистости",
    "performance": "общая эффективность",
    "scalp_health_comfort": "здоровье и комфорт кожи головы",
    "hair_feel": "ощущение волос",
    "price_affordability": "цена и доступность",
    "value_for_money": "соотношение цены и качества",
    "volume": "объём волос",
    "residue_buildup": "остаток и накопление продукта",
    "manageability": "послушность волос",
    "greasy_oily_finish": "жирность после применения",
    "ingredients_formula": "состав и формула",
    "hair_strength_breakage": "прочность и ломкость волос",
    "damage_repair_hair_health": "восстановление и здоровье волос",
    "dispenser_functionality": "работа дозатора",
    "hair_growth_loss_shedding": "рост и выпадение волос",
    "curls": "кудри и волны",
    "product_size": "размер и объём продукта",
    "application_and_usability": "нанесение и удобство",
    "conditioning_effect": "кондиционирующий эффект",
    "color_protection": "защита цвета",
    "use_case": "сценарий использования",
    "skin_sensitivity_reactions": "чувствительность и реакции кожи",
    "cleaning": "очищение",
    "longevity": "длительность эффекта",
    "heat_protection": "термозащита",
    "lather": "пена",
    "color_toning": "тонирование цвета",
    "slip": "скольжение продукта",
    "hold": "фиксация",
    "rinsing": "смывание",
}


def _evidence_text(value: str) -> str:
    return " | ".join(str(item["text"]) for item in json.loads(value))


def _diverse_take(
    candidates: list[dict[str, Any]],
    count: int,
    used_reviews: set[str],
) -> list[dict[str, Any]]:
    selected: list[dict[str, Any]] = []
    aspect_counts: dict[str, int] = {}
    remaining = list(candidates)
    while remaining and len(selected) < count:
        remaining.sort(
            key=lambda item: (
                item["review_id"] in used_reviews,
                aspect_counts.get(item["aspect_id"], 0),
                item["sort_key"],
            )
        )
        item = remaining.pop(0)
        selected.append(item)
        used_reviews.add(item["review_id"])
        aspect_counts[item["aspect_id"]] = aspect_counts.get(item["aspect_id"], 0) + 1
    return selected


def build_aspect_audit_workbook(
    sample_path: str | Path,
    prediction_path: str | Path,
    sentiment_path: str | Path,
    silver_path: str | Path,
    output_path: str | Path,
    *,
    rows_per_case: int = 8,
    report_path: str | Path | None = None,
) -> dict[str, Any]:
    """Select diverse high-value disagreements and write a Russian workbook."""
    sample = pd.read_parquet(sample_path)[
        ["review_id", "review_text", "sampling_component", "annotation_order"]
    ]
    predictions = pd.read_parquet(prediction_path)
    sentiments = pd.read_parquet(sentiment_path)
    silver = pd.read_parquet(silver_path)
    sample_by_id = sample.set_index("review_id").to_dict("index")
    prediction_by_pair = {
        (str(row.review_id), str(row.aspect_id)): row
        for row in predictions.itertuples(index=False)
    }
    sentiment_by_pair = {
        (str(row.review_id), str(row.aspect_id)): str(row.aspect_sentiment)
        for row in sentiments.itertuples(index=False)
    }
    silver_by_pair: dict[tuple[str, str], dict[str, Any]] = {}
    for row in silver.itertuples(index=False):
        for aspect in json.loads(row.aspects_json):
            silver_by_pair[(str(row.review_id), str(aspect["aspect_id"]))] = aspect

    predicted_pairs = set(prediction_by_pair)
    silver_pairs = set(silver_by_pair)
    false_positive: list[dict[str, Any]] = []
    for pair in predicted_pairs - silver_pairs:
        review_id, aspect_id = pair
        row = prediction_by_pair[pair]
        context = sample_by_id[review_id]
        false_positive.append(
            {
                "case_type": "лишний аспект модели",
                "review_id": review_id,
                "aspect_id": aspect_id,
                "model_evidence": _evidence_text(row.evidence_json),
                "silver_evidence": "",
                "model_probability": float(row.aspect_probability),
                "model_sentiment": sentiment_by_pair.get(pair, ""),
                "silver_sentiment": "",
                "sort_key": (
                    0 if context["sampling_component"] == "representative" else 1,
                    -float(row.aspect_probability),
                    int(context["annotation_order"]),
                ),
            }
        )
    false_negative: list[dict[str, Any]] = []
    for pair in silver_pairs - predicted_pairs:
        review_id, aspect_id = pair
        context = sample_by_id[review_id]
        aspect = silver_by_pair[pair]
        false_negative.append(
            {
                "case_type": "пропущенный аспект",
                "review_id": review_id,
                "aspect_id": aspect_id,
                "model_evidence": "",
                "silver_evidence": " | ".join(aspect["customer_phrases"]),
                "model_probability": None,
                "model_sentiment": "",
                "silver_sentiment": str(aspect["sentiment"]),
                "sort_key": (
                    0 if context["sampling_component"] == "representative" else 1,
                    int(context["annotation_order"]),
                ),
            }
        )
    sentiment_disagreement: list[dict[str, Any]] = []
    for pair in predicted_pairs & silver_pairs:
        predicted_sentiment = sentiment_by_pair.get(pair)
        silver_sentiment = str(silver_by_pair[pair]["sentiment"])
        if predicted_sentiment is None or predicted_sentiment == silver_sentiment:
            continue
        review_id, aspect_id = pair
        row = prediction_by_pair[pair]
        context = sample_by_id[review_id]
        sentiment_disagreement.append(
            {
                "case_type": "разная тональность",
                "review_id": review_id,
                "aspect_id": aspect_id,
                "model_evidence": _evidence_text(row.evidence_json),
                "silver_evidence": " | ".join(
                    silver_by_pair[pair]["customer_phrases"]
                ),
                "model_probability": float(row.aspect_probability),
                "model_sentiment": predicted_sentiment,
                "silver_sentiment": silver_sentiment,
                "sort_key": (
                    0 if context["sampling_component"] == "representative" else 1,
                    -float(row.aspect_probability),
                    int(context["annotation_order"]),
                ),
            }
        )

    used_reviews: set[str] = set()
    selected = []
    for candidates in (false_positive, false_negative, sentiment_disagreement):
        selected.extend(_diverse_take(candidates, rows_per_case, used_reviews))
    rows = []
    for number, item in enumerate(selected, start=1):
        context = sample_by_id[item["review_id"]]
        rows.append(
            {
                "№": number,
                "тип проверки": item["case_type"],
                "выборка": context["sampling_component"],
                "текст отзыва (англ.)": context["review_text"],
                "аспект ID": item["aspect_id"],
                "аспект по-русски": ASPECT_NAMES_RU.get(item["aspect_id"], ""),
                "доказательство модели": item["model_evidence"],
                "доказательство Gemini": item["silver_evidence"],
                "вероятность модели": item["model_probability"],
                "тональность модели": item["model_sentiment"],
                "тональность Gemini": item["silver_sentiment"],
                "аспект есть? yes/no": "",
                "правильная тональность": "",
                "комментарий": "",
            }
        )
    frame = pd.DataFrame(rows)
    destination = Path(output_path)
    destination.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(destination, engine="openpyxl") as writer:
        instructions = pd.DataFrame(
            {
                "Простая инструкция": [
                    "Проверьте только выделенный аспект по тексту отзыва.",
                    "В колонке «аспект есть?» выберите yes или no.",
                    "Если аспект есть, выберите его тональность; иначе оставьте пусто.",
                    "24 строки достаточно: это аудит разногласий, а не новая большая разметка.",
                ]
            }
        )
        escape_dataframe_for_spreadsheet(instructions).to_excel(
            writer, sheet_name="Инструкция", index=False
        )
        escape_dataframe_for_spreadsheet(frame).to_excel(
            writer, sheet_name="Проверка", index=False
        )
        sheet = writer.book["Проверка"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = Font(color="FFFFFF", bold=True)
            cell.alignment = Alignment(wrap_text=True, vertical="top")
        widths = [6, 24, 16, 70, 30, 30, 55, 55, 18, 22, 22, 20, 24, 35]
        for column, width in zip(sheet.columns, widths):
            sheet.column_dimensions[column[0].column_letter].width = width
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")
        yes_no = DataValidation(type="list", formula1='"yes,no"')
        tones = DataValidation(
            type="list", formula1='"negative,neutral,positive,mixed,unclear"'
        )
        sheet.add_data_validation(yes_no)
        sheet.add_data_validation(tones)
        yes_no.add(f"L2:L{len(frame) + 1}")
        tones.add(f"M2:M{len(frame) + 1}")
        instruction_sheet = writer.book["Инструкция"]
        instruction_sheet.column_dimensions["A"].width = 100
        for row in instruction_sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    report = {
        "audit_kind": "human_check_of_model_vs_gemini_silver_disagreements",
        "row_count": len(frame),
        "rows_per_case": rows_per_case,
        "case_counts": {
            str(name): int(value)
            for name, value in frame["тип проверки"].value_counts().sort_index().items()
        },
        "output_path": str(destination),
        "output_sha256": sha256_file(destination),
        "full_human_gold_required": False,
    }
    if report_path is not None:
        Path(report_path).write_text(
            json.dumps(report, ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("sample_path", type=Path)
    parser.add_argument("prediction_path", type=Path)
    parser.add_argument("sentiment_path", type=Path)
    parser.add_argument("silver_path", type=Path)
    parser.add_argument("output_path", type=Path)
    parser.add_argument("--rows-per-case", type=int, default=8)
    parser.add_argument("--report-path", type=Path)
    args = parser.parse_args()
    report = build_aspect_audit_workbook(
        args.sample_path,
        args.prediction_path,
        args.sentiment_path,
        args.silver_path,
        args.output_path,
        rows_per_case=args.rows_per_case,
        report_path=args.report_path,
    )
    print(json.dumps(report, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
