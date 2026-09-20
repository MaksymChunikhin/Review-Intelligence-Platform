"""Build a reproducible seller-versus-competitor demo report."""

from __future__ import annotations

import argparse
import json
from dataclasses import asdict, dataclass
from pathlib import Path
from typing import Any
from uuid import uuid4

import numpy as np
import pandas as pd
import duckdb
from openpyxl.styles import Alignment, Font, PatternFill

from src.common.csv_safety import escape_dataframe_for_spreadsheet
from src.common.project import find_project_root
from src.ingestion.dataset_manifest import sha256_file
from src.schemas.workspace import load_seller_workspace_scope


@dataclass(frozen=True)
class SellerComparisonReport:
    report_version: str
    workspace_id: str
    workspace_version: str
    seller_parent_asins: list[str]
    competitor_parent_asins: list[str]
    seller_review_count: int
    competitor_review_count: int
    comparison_aspect_count: int
    minimum_aspect_support: int
    minimum_comparison_gap_pp: float
    workbook_path: str
    workbook_sha256: str
    markdown_path: str
    markdown_sha256: str
    comparison_path: str
    comparison_sha256: str
    warning: str


def _artifact_reference(path: Path) -> str:
    try:
        return path.resolve().relative_to(find_project_root(path.parent)).as_posix()
    except (FileNotFoundError, ValueError):
        return str(path)


def _write_parquet(frame: pd.DataFrame, destination: Path) -> None:
    destination.parent.mkdir(parents=True, exist_ok=True)
    temporary = destination.with_name(
        f".{destination.stem}.{uuid4().hex}.tmp{destination.suffix}"
    )
    try:
        frame.to_parquet(temporary, index=False)
        temporary.replace(destination)
    finally:
        temporary.unlink(missing_ok=True)


def _read_products(path: Path, parent_asins: list[str]) -> pd.DataFrame:
    """Read only workspace products from a potentially large Parquet file."""
    connection = duckdb.connect()
    try:
        connection.execute("SET memory_limit='256MB'")
        connection.execute("SET threads=2")
        return connection.execute(
            "SELECT * FROM read_parquet(?) "
            "WHERE parent_asin IN (SELECT UNNEST(?))",
            [str(path), parent_asins],
        ).fetchdf()
    finally:
        connection.close()


def _aggregate_aspects(
    product_mart: pd.DataFrame,
    parent_asins: list[str],
    review_denominator: int,
    prefix: str,
) -> pd.DataFrame:
    selected = product_mart[product_mart["parent_asin"].isin(parent_asins)]
    result = selected.groupby("aspect_id", sort=True).agg(
        aspect_review_count=("aspect_review_count", "sum"),
        needs_attention_count=("needs_attention_count", "sum"),
        satisfied_count=("satisfied_count", "sum"),
        negative_count=("negative_count", "sum"),
        neutral_count=("neutral_count", "sum"),
        positive_count=("positive_count", "sum"),
        product_count_with_aspect=("parent_asin", "nunique"),
    )
    denominator = result["aspect_review_count"].replace(0, np.nan)
    result["aspect_review_share"] = result["aspect_review_count"] / review_denominator
    result["needs_attention_share"] = result["needs_attention_count"] / denominator
    result["satisfied_share"] = result["satisfied_count"] / denominator
    result["negative_share"] = result["negative_count"] / denominator
    result["neutral_share"] = result["neutral_count"] / denominator
    result["positive_share"] = result["positive_count"] / denominator
    return result.add_prefix(f"{prefix}_").reset_index()


def _product_overview(
    reviews: pd.DataFrame,
    product_mart: pd.DataFrame,
    seller_ids: list[str],
    competitor_ids: list[str],
) -> pd.DataFrame:
    selected_ids = seller_ids + competitor_ids
    selected = reviews[reviews["parent_asin"].isin(selected_ids)].copy()
    selected["needs_attention"] = selected["rating"] <= 3
    selected["year_month"] = pd.to_datetime(
        selected["review_timestamp"], utc=True
    ).dt.strftime("%Y-%m")
    overview = selected.groupby("parent_asin", sort=False).agg(
        product_title=("product_title", "first"),
        store=("store", "first"),
        review_count=("review_id", "size"),
        average_rating=("rating", "mean"),
        needs_attention_count=("needs_attention", "sum"),
        first_review=("review_timestamp", "min"),
        last_review=("review_timestamp", "max"),
        active_months=("year_month", "nunique"),
        helpful_votes=("helpful_vote", "sum"),
    ).reset_index()
    overview["needs_attention_share"] = (
        overview["needs_attention_count"] / overview["review_count"]
    )
    aspect_counts = product_mart[
        product_mart["parent_asin"].isin(selected_ids)
    ].groupby("parent_asin").agg(
        aspect_observation_count=("aspect_review_count", "sum"),
        distinct_aspect_count=("aspect_id", "nunique"),
    ).reset_index()
    overview = overview.merge(
        aspect_counts, on="parent_asin", how="left", validate="one_to_one"
    )
    overview["role"] = overview["parent_asin"].map(
        lambda value: "seller" if value in seller_ids else "competitor"
    )
    for column in ("first_review", "last_review"):
        overview[column] = pd.to_datetime(overview[column], utc=True).dt.strftime(
            "%Y-%m-%d"
        )
    overview["display_order"] = overview["parent_asin"].map(
        {value: index for index, value in enumerate(selected_ids)}
    )
    return overview.sort_values("display_order").drop(columns="display_order")


def _format_percent(value: float) -> str:
    return "—" if pd.isna(value) else f"{value * 100:.1f}%"


def _markdown_table(frame: pd.DataFrame, columns: list[str]) -> str:
    header = "| " + " | ".join(columns) + " |"
    divider = "|" + "|".join(["---"] * len(columns)) + "|"
    rows = [
        "| " + " | ".join(str(row[column]) for column in columns) + " |"
        for _, row in frame.iterrows()
    ]
    return "\n".join([header, divider, *rows])


def build_seller_comparison(
    workspace_path: str | Path,
    reviews_path: str | Path,
    product_aspect_mart_path: str | Path,
    comparison_path: str | Path,
    workbook_path: str | Path,
    markdown_path: str | Path,
    *,
    report_version: str,
    minimum_aspect_support: int = 20,
    minimum_comparison_gap_pp: float = 5.0,
    report_path: str | Path | None = None,
) -> SellerComparisonReport:
    """Compare one seller scope with a review-weighted competitor benchmark."""
    if minimum_aspect_support < 1:
        raise ValueError("minimum_aspect_support must be positive")
    if minimum_comparison_gap_pp < 0:
        raise ValueError("minimum_comparison_gap_pp must not be negative")
    workspace = load_seller_workspace_scope(workspace_path)
    reviews_source = Path(reviews_path)
    product_source = Path(product_aspect_mart_path)
    comparison_destination = Path(comparison_path)
    workbook_destination = Path(workbook_path)
    markdown_destination = Path(markdown_path)
    selected_ids = workspace.seller_parent_asins + workspace.competitor_parent_asins
    reviews = _read_products(reviews_source, selected_ids)
    product_mart = _read_products(product_source, selected_ids)
    missing = set(selected_ids) - set(reviews["parent_asin"].astype(str))
    if missing:
        raise ValueError(f"workspace products are absent from review population: {missing}")
    start = pd.Timestamp(workspace.review_date_start, tz="UTC")
    end = pd.Timestamp(workspace.review_date_end, tz="UTC") + pd.Timedelta(days=1)
    review_time = pd.to_datetime(reviews["review_timestamp"], utc=True)
    reviews = reviews[
        reviews["parent_asin"].isin(selected_ids)
        & (review_time >= start)
        & (review_time < end)
    ].copy()
    if workspace.verified_purchase_only:
        reviews = reviews[reviews["verified_purchase"]].copy()
    valid_review_ids = set(reviews["review_id"].astype(str))
    product_mart = product_mart[
        product_mart["parent_asin"].isin(selected_ids)
    ].copy()
    # Comparative claims are allowed only for aspects that passed the
    # reliability gate. Legacy marts do not have this field and retain their
    # historical behavior.
    if "comparative_analytics_allowed" in product_mart.columns:
        product_mart = product_mart[
            product_mart["comparative_analytics_allowed"].astype(bool)
        ].copy()
    # Current marts use the same full date scope. Fail instead of silently
    # mixing a narrower workspace with pre-aggregated counts.
    if start.date() != pd.Timestamp("2021-01-01").date() or end.date() != pd.Timestamp(
        "2023-09-13"
    ).date():
        raise ValueError("current product mart supports the full snapshot date scope only")
    if workspace.verified_purchase_only:
        raise ValueError("current product mart does not support verified-only workspaces")
    if not valid_review_ids:
        raise ValueError("workspace resolves to no reviews")

    overview = _product_overview(
        reviews,
        product_mart,
        workspace.seller_parent_asins,
        workspace.competitor_parent_asins,
    )
    seller_review_count = int(
        overview.loc[overview["role"] == "seller", "review_count"].sum()
    )
    competitor_review_count = int(
        overview.loc[overview["role"] == "competitor", "review_count"].sum()
    )
    seller = _aggregate_aspects(
        product_mart,
        workspace.seller_parent_asins,
        seller_review_count,
        "seller",
    )
    competitor = _aggregate_aspects(
        product_mart,
        workspace.competitor_parent_asins,
        competitor_review_count,
        "competitor",
    )
    aspect_definitions = product_mart[
        ["aspect_id", "canonical_name", "parent_group"]
    ].drop_duplicates("aspect_id")
    comparison = aspect_definitions.merge(
        seller, on="aspect_id", how="left", validate="one_to_one"
    ).merge(competitor, on="aspect_id", how="left", validate="one_to_one")
    count_columns = [
        column
        for column in comparison.columns
        if column.endswith("_count") or column.endswith("_count_with_aspect")
    ]
    comparison[count_columns] = comparison[count_columns].fillna(0).astype(int)
    comparison["discussion_gap_pp"] = 100 * (
        comparison["seller_aspect_review_share"]
        - comparison["competitor_aspect_review_share"]
    )
    comparison["attention_gap_pp"] = 100 * (
        comparison["seller_needs_attention_share"]
        - comparison["competitor_needs_attention_share"]
    )
    comparison["positive_gap_pp"] = 100 * (
        comparison["seller_positive_share"]
        - comparison["competitor_positive_share"]
    )
    comparison["advantage_score_pp"] = (
        comparison["positive_gap_pp"] - comparison["attention_gap_pp"]
    )
    comparison["support_sufficient"] = (
        comparison["seller_aspect_review_count"] >= minimum_aspect_support
    ) & (
        comparison["competitor_aspect_review_count"] >= minimum_aspect_support
    )
    comparison = comparison.sort_values(
        ["seller_aspect_review_count", "aspect_id"], ascending=[False, True]
    ).reset_index(drop=True)
    _write_parquet(comparison, comparison_destination)

    supported = comparison[comparison["support_sufficient"]].copy()
    complaints = supported.sort_values(
        ["seller_needs_attention_count", "attention_gap_pp"], ascending=False
    ).head(10)
    weaknesses = supported[
        supported["attention_gap_pp"] >= minimum_comparison_gap_pp
    ].sort_values(
        ["attention_gap_pp", "seller_needs_attention_count"], ascending=False
    ).head(10)
    strengths = supported[
        (supported["seller_positive_share"] >= 0.55)
        & (supported["positive_gap_pp"] >= minimum_comparison_gap_pp)
        & (supported["attention_gap_pp"] <= -minimum_comparison_gap_pp)
    ].sort_values(
        ["advantage_score_pp", "seller_positive_count"], ascending=False
    ).head(10)

    seller_evidence = product_mart[
        product_mart["parent_asin"].isin(workspace.seller_parent_asins)
    ][
        [
            "aspect_id",
            "top_negative_review_id",
            "top_negative_evidence_json",
            "top_positive_review_id",
            "top_positive_evidence_json",
        ]
    ].copy()
    aspect_name_by_id = aspect_definitions.set_index("aspect_id")[
        "canonical_name"
    ].to_dict()
    seller_evidence["aspect_name"] = seller_evidence["aspect_id"].map(
        aspect_name_by_id
    )
    def readable_evidence(value: Any) -> str:
        if pd.isna(value):
            return ""
        return " | ".join(str(item["text"]) for item in json.loads(str(value)))

    seller_evidence["negative_evidence"] = seller_evidence[
        "top_negative_evidence_json"
    ].map(readable_evidence)
    seller_evidence["positive_evidence"] = seller_evidence[
        "top_positive_evidence_json"
    ].map(readable_evidence)
    seller_evidence = seller_evidence[
        [
            "aspect_id",
            "aspect_name",
            "top_negative_review_id",
            "negative_evidence",
            "top_positive_review_id",
            "positive_evidence",
        ]
    ]
    negative_evidence_by_aspect = seller_evidence.set_index("aspect_id")[
        "negative_evidence"
    ].to_dict()
    positive_evidence_by_aspect = seller_evidence.set_index("aspect_id")[
        "positive_evidence"
    ].to_dict()

    workbook_destination.parent.mkdir(parents=True, exist_ok=True)
    export_frames = {
        "Products": overview,
        "Aspect comparison": comparison,
        "Top complaints": complaints,
        "Weaknesses": weaknesses,
        "Strengths": strengths,
        "Review evidence": seller_evidence,
    }
    with pd.ExcelWriter(workbook_destination, engine="openpyxl") as writer:
        for sheet_name, frame in export_frames.items():
            escape_dataframe_for_spreadsheet(frame).to_excel(
                writer, sheet_name=sheet_name, index=False
            )
        for sheet in writer.book.worksheets:
            sheet.freeze_panes = "A2"
            sheet.auto_filter.ref = sheet.dimensions
            for cell in sheet[1]:
                cell.fill = PatternFill("solid", fgColor="1F4E78")
                cell.font = Font(color="FFFFFF", bold=True)
                cell.alignment = Alignment(wrap_text=True, vertical="top")
            for column in sheet.columns:
                values = [str(cell.value or "") for cell in column[:50]]
                width = min(55, max(12, max(map(len, values), default=12) + 2))
                sheet.column_dimensions[column[0].column_letter].width = width
            for row in sheet.iter_rows(min_row=2):
                for cell in row:
                    cell.alignment = Alignment(wrap_text=True, vertical="top")

    seller_row = overview[overview["role"] == "seller"].iloc[0]
    overview_display = overview[
        ["role", "parent_asin", "store", "review_count", "average_rating", "needs_attention_share"]
    ].copy()
    overview_display["average_rating"] = overview_display["average_rating"].map(
        lambda value: f"{value:.2f}"
    )
    overview_display["needs_attention_share"] = overview_display[
        "needs_attention_share"
    ].map(_format_percent)
    overview_display.columns = [
        "Role", "Parent ASIN", "Brand", "Reviews", "Rating", "1–3 star share"
    ]

    def insight_table(frame: pd.DataFrame, kind: str) -> pd.DataFrame:
        result = frame[
            [
                "aspect_id",
                "canonical_name",
                "seller_aspect_review_count",
                "seller_needs_attention_share",
                "competitor_needs_attention_share",
                "attention_gap_pp",
                "seller_positive_share",
                "competitor_positive_share",
            ]
        ].head(5).copy()
        if kind == "strengths":
            result["evidence"] = result["aspect_id"].map(
                positive_evidence_by_aspect
            )
        else:
            result["evidence"] = result["aspect_id"].map(
                negative_evidence_by_aspect
            )
        result = result.drop(columns="aspect_id")
        for column in (
            "seller_needs_attention_share",
            "competitor_needs_attention_share",
            "seller_positive_share",
            "competitor_positive_share",
        ):
            result[column] = result[column].map(_format_percent)
        result["attention_gap_pp"] = result["attention_gap_pp"].map(
            lambda value: "—" if pd.isna(value) else f"{value:+.1f} pp"
        )
        result.columns = [
            "Aspect", "Mentions", "Product 1–3 star share", "Competitor 1–3 star share",
            "Difference", "Product positive share", "Competitor positive share", "Review example",
        ]
        return result

    markdown = f"""# Seller comparison report: {seller_row['store']}

Product: `{seller_row['parent_asin']}` — {seller_row['product_title']}

Data period: 2021-01-01 — 2023-09-12. This is a historical local snapshot,
not a current Amazon listing. Aspect extraction and aspect sentiment are
diagnostic model signals; the 1–3 star signal comes directly from review ratings.

## Products in this comparison

{_markdown_table(overview_display, list(overview_display.columns))}

## Most-mentioned complaints

{_markdown_table(insight_table(complaints, 'complaints'), list(insight_table(complaints, 'complaints').columns))}

## Supported weaknesses

{_markdown_table(insight_table(weaknesses, 'weaknesses'), list(insight_table(weaknesses, 'weaknesses').columns))}

## Supported strengths

{_markdown_table(insight_table(strengths, 'strengths'), list(insight_table(strengths, 'strengths').columns))}

## Interpretation

- Improvement priorities combine aspect sentiment with the reliable 1–3 star signal.
- Competitors come from the same exact Amazon product path and exclude the same brand.
- Only aspects approved for comparative analytics are used in strengths and weaknesses.
- A strength or weakness requires at least {minimum_comparison_gap_pp:.1f} percentage points of separation.
- Full evidence excerpts are available on the `Review evidence` Excel sheet.
"""
    markdown_destination.parent.mkdir(parents=True, exist_ok=True)
    markdown_destination.write_text(markdown, encoding="utf-8")

    report = SellerComparisonReport(
        report_version=report_version,
        workspace_id=workspace.workspace_id,
        workspace_version=workspace.workspace_version,
        seller_parent_asins=workspace.seller_parent_asins,
        competitor_parent_asins=workspace.competitor_parent_asins,
        seller_review_count=seller_review_count,
        competitor_review_count=competitor_review_count,
        comparison_aspect_count=len(comparison),
        minimum_aspect_support=minimum_aspect_support,
        minimum_comparison_gap_pp=minimum_comparison_gap_pp,
        workbook_path=_artifact_reference(workbook_destination),
        workbook_sha256=sha256_file(workbook_destination),
        markdown_path=_artifact_reference(markdown_destination),
        markdown_sha256=sha256_file(markdown_destination),
        comparison_path=_artifact_reference(comparison_destination),
        comparison_sha256=sha256_file(comparison_destination),
        warning=(
            "Historical demo using candidate aspect extraction and weak-label "
            "aspect sentiment; not current Amazon market data."
        ),
    )
    if report_path is not None:
        Path(report_path).write_text(
            json.dumps(asdict(report), ensure_ascii=False, indent=2) + "\n",
            encoding="utf-8",
        )
    return report


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("workspace_path", type=Path)
    parser.add_argument("reviews_path", type=Path)
    parser.add_argument("product_aspect_mart_path", type=Path)
    parser.add_argument("comparison_path", type=Path)
    parser.add_argument("workbook_path", type=Path)
    parser.add_argument("markdown_path", type=Path)
    parser.add_argument("--report-version", required=True)
    parser.add_argument("--minimum-aspect-support", type=int, default=20)
    parser.add_argument("--minimum-comparison-gap-pp", type=float, default=5.0)
    parser.add_argument("--report-path", type=Path)
    args = parser.parse_args()
    report = build_seller_comparison(
        args.workspace_path,
        args.reviews_path,
        args.product_aspect_mart_path,
        args.comparison_path,
        args.workbook_path,
        args.markdown_path,
        report_version=args.report_version,
        minimum_aspect_support=args.minimum_aspect_support,
        minimum_comparison_gap_pp=args.minimum_comparison_gap_pp,
        report_path=args.report_path,
    )
    print(json.dumps(asdict(report), ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
