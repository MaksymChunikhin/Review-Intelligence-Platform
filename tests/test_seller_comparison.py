"""Tests for seller-versus-competitor demo reporting."""

from __future__ import annotations

import json
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from src.analytics.seller_comparison import build_seller_comparison


def test_build_seller_comparison_writes_reproducible_outputs(tmp_path: Path) -> None:
    workspace_path = tmp_path / "workspace.json"
    workspace_path.write_text(
        json.dumps(
            {
                "workspace_id": "demo",
                "workspace_version": "demo_v1",
                "dataset_version": "dataset_v1",
                "niche": {
                    "niche_id": "conditioners",
                    "niche_version": "conditioners_v1",
                    "dataset_version": "dataset_v1",
                    "category_registry_schema_version": "registry_v1",
                    "display_name": "Conditioners",
                    "status": "approved",
                    "category_path_ids": ["path-1"],
                },
                "seller_parent_asins": ["p1"],
                "competitor_parent_asins": ["p2"],
                "review_date_start": "2021-01-01",
                "review_date_end": "2023-09-12",
                "verified_purchase_only": False,
            }
        ),
        encoding="utf-8",
    )
    reviews_path = tmp_path / "reviews.parquet"
    pd.DataFrame(
        {
            "review_id": ["r1", "r2"],
            "parent_asin": ["p1", "p2"],
            "rating": [2.0, 5.0],
            "review_timestamp": pd.to_datetime(
                ["2022-01-01", "2022-01-02"], utc=True
            ),
            "helpful_vote": [2, 1],
            "product_title": ["Seller product", "Competitor product"],
            "store": ["Seller", "Competitor"],
            "verified_purchase": [True, True],
        }
    ).to_parquet(reviews_path, index=False)
    product_path = tmp_path / "product_mart.parquet"
    pd.DataFrame(
        {
            "parent_asin": ["p1", "p2"],
            "aspect_id": ["scent", "scent"],
            "canonical_name": ["Scent", "Scent"],
            "parent_group": ["sensory", "sensory"],
            "aspect_review_count": [1, 1],
            "needs_attention_count": [1, 0],
            "satisfied_count": [0, 1],
            "negative_count": [1, 0],
            "neutral_count": [0, 0],
            "positive_count": [0, 1],
            "top_negative_review_id": ["r1", None],
            "top_negative_evidence_json": ['[{"text":"bad scent"}]', None],
            "top_positive_review_id": [None, "r2"],
            "top_positive_evidence_json": [None, '[{"text":"nice scent"}]'],
        }
    ).to_parquet(product_path, index=False)
    comparison_path = tmp_path / "comparison.parquet"
    workbook_path = tmp_path / "report.xlsx"
    markdown_path = tmp_path / "report.md"
    report = build_seller_comparison(
        workspace_path,
        reviews_path,
        product_path,
        comparison_path,
        workbook_path,
        markdown_path,
        report_version="report_v1",
        minimum_aspect_support=1,
    )
    comparison = pd.read_parquet(comparison_path)
    assert report.seller_review_count == 1
    assert report.competitor_review_count == 1
    assert comparison.iloc[0]["attention_gap_pp"] == 100.0
    assert "Marc" not in markdown_path.read_text(encoding="utf-8")
    assert "Weaknesses" in load_workbook(workbook_path).sheetnames


def test_seller_comparison_excludes_unreliable_aspects(tmp_path: Path) -> None:
    workspace = {
        "workspace_id": "demo",
        "workspace_version": "demo_v1",
        "dataset_version": "dataset_v1",
        "niche": {
            "niche_id": "products",
            "niche_version": "products_v1",
            "dataset_version": "dataset_v1",
            "category_registry_schema_version": "registry_v1",
            "display_name": "Products",
            "status": "approved",
            "category_path_ids": ["path-1"],
        },
        "seller_parent_asins": ["p1"],
        "competitor_parent_asins": ["p2"],
        "review_date_start": "2021-01-01",
        "review_date_end": "2023-09-12",
    }
    workspace_path = tmp_path / "workspace.json"
    workspace_path.write_text(json.dumps(workspace), encoding="utf-8")
    reviews = pd.DataFrame(
        {
            "review_id": ["r1", "r2"],
            "parent_asin": ["p1", "p2"],
            "rating": [5.0, 1.0],
            "review_timestamp": pd.to_datetime(["2022-01-01", "2022-01-02"], utc=True),
            "helpful_vote": [0, 0],
            "product_title": ["=unsafe", "Safe"],
            "store": ["Seller", "Competitor"],
            "verified_purchase": [True, True],
        }
    )
    reviews_path = tmp_path / "reviews.parquet"
    reviews.to_parquet(reviews_path, index=False)
    mart = pd.DataFrame(
        {
            "parent_asin": ["p1", "p2"],
            "aspect_id": ["experimental", "experimental"],
            "canonical_name": ["Experimental", "Experimental"],
            "parent_group": ["other", "other"],
            "aspect_review_count": [20, 20],
            "needs_attention_count": [0, 20],
            "satisfied_count": [20, 0],
            "negative_count": [0, 20],
            "neutral_count": [0, 0],
            "positive_count": [20, 0],
            "top_negative_review_id": [None, "r2"],
            "top_negative_evidence_json": [None, '[{"text":"bad"}]'],
            "top_positive_review_id": ["r1", None],
            "top_positive_evidence_json": ['[{"text":"good"}]', None],
            "comparative_analytics_allowed": [False, False],
        }
    )
    mart_path = tmp_path / "mart.parquet"
    mart.to_parquet(mart_path, index=False)
    comparison_path = tmp_path / "comparison.parquet"
    workbook_path = tmp_path / "report.xlsx"
    build_seller_comparison(
        workspace_path,
        reviews_path,
        mart_path,
        comparison_path,
        workbook_path,
        tmp_path / "report.md",
        report_version="report_v1",
    )
    assert pd.read_parquet(comparison_path).empty
    workbook = load_workbook(workbook_path, data_only=False)
    assert workbook["Products"]["B2"].value == "'=unsafe"
